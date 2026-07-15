"""Aggregate member rows and build Excel workbooks (Beam / Columns / Base Plate)."""

from __future__ import annotations

import io
from collections import defaultdict
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEET_COLUMNS = ["Member Name", "Quantity", "Size"]

# Extra detail columns kept after the required three (for tracing completeness)
BEAM_EXTRA = ["Length", "Material", "Page"]
COLUMN_EXTRA = ["Height", "Material", "Page"]
PLATE_EXTRA = ["Thickness", "Anchor Bolt Dia", "Anchor Bolt Qty", "Top of Concrete EL", "Page"]


def _merge_key(row: dict[str, Any]) -> tuple[str, str]:
    name = str(row.get("Member Name") or row.get("Mark") or "").strip().upper()
    size = str(row.get("Size") or row.get("Section Size") or row.get("Plate Size") or "").strip().upper()
    return name, size


def aggregate_members(
    rows: list[dict[str, Any]],
    extra_keys: list[str] | None = None,
) -> list[dict[str, Any]]:
    """
    Collapse duplicate marks into one row with summed Quantity.

    Same Member Name + Size → quantity added. First non-empty extras kept.
    """
    extra_keys = extra_keys or []
    buckets: dict[tuple[str, str], dict[str, Any]] = {}
    order: list[tuple[str, str]] = []

    for row in rows:
        key = _merge_key(row)
        if not key[0]:
            continue
        qty = row.get("Quantity", 1)
        try:
            qty_i = int(qty)
        except (TypeError, ValueError):
            qty_i = 1
        if qty_i < 1:
            qty_i = 1

        if key not in buckets:
            buckets[key] = {
                "Member Name": key[0],
                "Quantity": qty_i,
                "Size": key[1],
            }
            for ek in extra_keys:
                buckets[key][ek] = row.get(ek, "") or ""
            order.append(key)
        else:
            buckets[key]["Quantity"] += qty_i
            for ek in extra_keys:
                if not buckets[key].get(ek) and row.get(ek):
                    buckets[key][ek] = row[ek]

    return [buckets[k] for k in order]


def _style_sheet(ws, header_fill: str) -> None:
    fill = PatternFill("solid", fgColor=header_fill)
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max(12, max_len + 3), 42)


def build_excel_bytes(
    beams: list[dict[str, Any]],
    columns: list[dict[str, Any]],
    base_plates: list[dict[str, Any]],
    *,
    include_details: bool = True,
) -> bytes:
    """
    Build an .xlsx workbook with three sheets:
      - Beam
      - Columns
      - Base Plate

    Each sheet always includes Member Name, Quantity, and Size.
    """
    agg_beams = aggregate_members(beams, BEAM_EXTRA if include_details else [])
    agg_cols = aggregate_members(columns, COLUMN_EXTRA if include_details else [])
    agg_plates = aggregate_members(base_plates, PLATE_EXTRA if include_details else [])

    beam_cols = SHEET_COLUMNS + (BEAM_EXTRA if include_details else [])
    col_cols = SHEET_COLUMNS + (COLUMN_EXTRA if include_details else [])
    plate_cols = SHEET_COLUMNS + (PLATE_EXTRA if include_details else [])

    df_beams = pd.DataFrame(agg_beams, columns=beam_cols) if agg_beams else pd.DataFrame(columns=beam_cols)
    df_cols = pd.DataFrame(agg_cols, columns=col_cols) if agg_cols else pd.DataFrame(columns=col_cols)
    df_plates = (
        pd.DataFrame(agg_plates, columns=plate_cols) if agg_plates else pd.DataFrame(columns=plate_cols)
    )

    # Ensure Quantity is int
    for df in (df_beams, df_cols, df_plates):
        if "Quantity" in df.columns and len(df):
            df["Quantity"] = df["Quantity"].fillna(1).astype(int)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_beams.to_excel(writer, sheet_name="Beam", index=False)
        df_cols.to_excel(writer, sheet_name="Columns", index=False)
        df_plates.to_excel(writer, sheet_name="Base Plate", index=False)

        _style_sheet(writer.sheets["Beam"], "1E3A5F")
        _style_sheet(writer.sheets["Columns"], "1F4E3D")
        _style_sheet(writer.sheets["Base Plate"], "6B3A1F")

    buffer.seek(0)
    return buffer.read()


def members_preview(
    beams: list[dict[str, Any]],
    columns: list[dict[str, Any]],
    base_plates: list[dict[str, Any]],
) -> dict[str, Any]:
    """JSON-friendly preview with aggregated required columns."""
    agg_b = aggregate_members(beams, BEAM_EXTRA)
    agg_c = aggregate_members(columns, COLUMN_EXTRA)
    agg_p = aggregate_members(base_plates, PLATE_EXTRA)
    return {
        "counts": {
            "beams": len(agg_b),
            "columns": len(agg_c),
            "base_plates": len(agg_p),
            "beam_quantity_total": sum(r["Quantity"] for r in agg_b),
            "column_quantity_total": sum(r["Quantity"] for r in agg_c),
            "base_plate_quantity_total": sum(r["Quantity"] for r in agg_p),
        },
        "Beam": [{k: r.get(k, "") for k in SHEET_COLUMNS} for r in agg_b],
        "Columns": [{k: r.get(k, "") for k in SHEET_COLUMNS} for r in agg_c],
        "Base Plate": [{k: r.get(k, "") for k in SHEET_COLUMNS} for r in agg_p],
    }


def count_by_size(rows: list[dict[str, Any]]) -> dict[str, int]:
    """Utility: total quantity grouped by Size only."""
    totals: dict[str, int] = defaultdict(int)
    for row in aggregate_members(rows):
        size = row.get("Size") or "UNKNOWN"
        totals[size] += int(row.get("Quantity") or 1)
    return dict(totals)
