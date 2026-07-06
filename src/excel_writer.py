"""Write the BOM workbook with a Detailed_BOM sheet and a Summary sheet."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

_WIDTHS = {
    "Member_ID": 14,
    "Member_Type": 14,
    "Section_Size": 20,
    "Length_mm": 12,
    "Quantity": 10,
    "Elevation_m": 12,
    "Grid_Location": 14,
    "Weight_kg": 12,
    "Source_File": 22,
    "Total_Qty": 12,
    "Total_Length_mm": 18,
    "Total_Weight_kg": 16,
    "Total_Count": 12,
}


def _autofit(worksheet, df: pd.DataFrame, start_col: int = 1) -> None:
    from openpyxl.utils import get_column_letter

    for offset, column in enumerate(df.columns):
        letter = get_column_letter(start_col + offset)
        worksheet.column_dimensions[letter].width = _WIDTHS.get(column, 15)


def write_bom(
    detailed: pd.DataFrame,
    summary: pd.DataFrame,
    type_counts: pd.DataFrame,
    output_path: Path,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        detailed.to_excel(writer, sheet_name="Detailed_BOM", index=False)
        _autofit(writer.sheets["Detailed_BOM"], detailed)

        # Summary sheet: section totals, then a member-type count block below.
        summary.to_excel(writer, sheet_name="Summary", index=False, startrow=1)
        sheet = writer.sheets["Summary"]
        sheet.cell(row=1, column=1, value="Section Summary (Total Qty & Length per Section_Size)")
        _autofit(sheet, summary)

        type_start = len(summary) + 4
        sheet.cell(row=type_start, column=1, value="Member Type Counts")
        type_counts.to_excel(writer, sheet_name="Summary", index=False, startrow=type_start)

    return output_path
